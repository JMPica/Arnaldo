# -*- coding: utf-8 -*-
"""Auditoría independiente del Financial Model: recalcula con LibreOffice los escenarios Base y Optimista
(cambiando Assumptions!C4), rehace VAN y TIR fuera del modelo y comprueba el reparto unitario del PVP.
Añade los resultados a metrics.json (claves audit_*). Uso: python3 audit_model.py ../PROYECTO_CIGARS_CUBA_2027 metrics.json"""
import json, os, subprocess, sys, tempfile
import openpyxl

R, MJ = sys.argv[1], sys.argv[2]
SRC = os.path.join(R, "03_BUSINESS_CASE/Financial Model 2027–2031.xlsx")
M = json.load(open(MJ))
T = tempfile.mkdtemp()


def es(x, d=0):
    return f"{x:,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")


def npv(r, cf):
    return sum(c / (1 + r) ** (t + 1) for t, c in enumerate(cf))


def irr(cf):
    lo, hi = -0.99, 5.0
    if npv(lo, cf) * npv(hi, cf) > 0:
        return None
    for _ in range(200):
        mid = (lo + hi) / 2
        lo, hi = (lo, mid) if npv(lo, cf) * npv(mid, cf) <= 0 else (mid, hi)
    return mid


def scenario(sc):
    wb = openpyxl.load_workbook(SRC)
    wb["Assumptions"]["C4"] = sc
    p = os.path.join(T, f"m{sc}.xlsx")
    wb.save(p)
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", os.path.join(T, "rc"), p],
                   check=True, capture_output=True)
    w = openpyxl.load_workbook(os.path.join(T, "rc", f"m{sc}.xlsx"), data_only=True)
    rows = {r[0]: r[1:6] for r in w["Investment Returns"].iter_rows(values_only=True) if r[0]}
    return [float(x) for x in rows["FCF del proyecto"]], [float(x) for x in rows["Valor terminal"]], \
        rows["VAN al coste de capital (€)"][0], rows["TIR del proyecto"][0]


checks, ok = [], True
rate = 0.18
for sc, k in [(2, "base"), (3, "opt")]:
    fcf, tv, npv_m, irr_m = scenario(sc)
    tot = [a + b for a, b in zip(fcf, tv)]
    n_ind = npv(rate, tot)
    checks.append((f"VAN {k}", abs(n_ind - npv_m) < 1))
    i_ind = irr(tot)
    if isinstance(irr_m, (int, float)):
        checks.append((f"TIR {k}", i_ind is not None and abs(i_ind - irr_m) < 1e-6))
    M[f"audit_{k}_npv_no_tv"] = es(npv(rate, fcf)) + " €"
    M[f"audit_{k}_pv_tv"] = es(sum(v / (1 + rate) ** (t + 1) for t, v in enumerate(tv))) + " €"

# Reparto unitario (escenario Base, 2028): IVA, IEE, canal y arancel recalculados desde las reglas
pvp = 24.0
iva, iee = pvp * 0.21 / 1.21, max(0.158 * pvp, 47 / 1000)
duty = 0.26 * M["split_cif_eur_raw"]
for key, val in (("vat", iva), ("iee", iee), ("rm", 0.09 * pvp), ("dist", 0.05 * pvp), ("duty", duty)):
    checks.append((f"reparto {key}", abs(val - M[f"split_{key}_eur_raw"]) < 1e-6))
total = sum(M[f"split_{k}_eur_raw"] for k in ("vat", "iee", "rm", "dist", "duty", "cif", "fiscal", "merma", "royalty", "contrib"))
checks.append(("reparto suma PVP", abs(total - pvp) < 1e-6))

ok = all(c[1] for c in checks)
M["audit_checks"] = f"{sum(c[1] for c in checks)}/{len(checks)}"
M["audit_ok"] = "OK" if ok else "ERROR"
json.dump(M, open(MJ, "w"), ensure_ascii=False, indent=1)
for name, good in checks:
    print(("OK   " if good else "FAIL ") + name)
sys.exit(0 if ok else 1)
