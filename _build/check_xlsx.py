import sys
from openpyxl import load_workbook
wb = load_workbook(sys.argv[1], data_only=True)
errs = 0
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and (v.startswith("#") or v.startswith("Err:")) and v not in ("#",):
                if errs < 40: print("ERR", ws.title, c.coordinate, v)
                errs += 1
print("total errors", errs)
for name in sys.argv[2:]:
    ws = wb[name]
    print("=====", name)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        vals = [c.value for c in row[:8]]
        if any(v is not None for v in vals):
            print(" | ".join("" if v is None else (f"{v:,.2f}" if isinstance(v,(int,float)) else str(v)[:55]) for v in vals))
