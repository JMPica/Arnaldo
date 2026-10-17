# -*- coding: utf-8 -*-
"""Financial Model 2027–2031 (Proyecto A — marca propia ARNALDO VALLÉS).

Arquitectura: un único motor de cálculo (hoja 'Engine') calcula los 4 escenarios
en bloques paralelos con las MISMAS fórmulas. Las hojas de presentación
(P&L, Cash Flow, Balance...) leen el escenario seleccionado en Assumptions!C4.
Todas las cifras derivadas son fórmulas; sólo los supuestos son valores.
"""
import sys
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L
from xstyle import *
from sources import SOURCES

OUT = sys.argv[1] if len(sys.argv) > 1 else "Financial Model 2027-2031.xlsx"
YEARS = [2027, 2028, 2029, 2030, 2031]
SCEN = ["Conservador", "Base", "Optimista", "Estrés"]
NY = 5

wb = Workbook()
names = ["Dashboard", "Assumptions", "Product Cost", "Import & Taxes", "Unit Economics",
         "Sales Forecast", "Operating Expenses", "Working Capital", "Profit & Loss", "Cash Flow",
         "Balance Sheet", "Break-even", "Scenarios", "Sensitivity Analysis", "Investment Returns",
         "Risk Register", "Sources & Validation", "Engine"]
ws_ = {}
for i, n in enumerate(names):
    ws_[n] = wb.active if i == 0 else wb.create_sheet(n)
    ws_[n].title = n


def q(sheet):
    return "'%s'" % sheet if (" " in sheet or "&" in sheet or "-" in sheet) else sheet


# ---------------------------------------------------------------- Assumptions
A = ws_["Assumptions"]
title(A, "Supuestos del modelo — Proyecto A (marca propia ARNALDO VALLÉS)",
      "Todos los valores en azul son editables. Estados: VERIFICADO / PROBABLE / HIPÓTESIS / BLOQUEADO. Fuentes en 'Sources & Validation'. Fecha de corte: 10/10/2026.")
legend(A, 3, 3)
put(A, "A4", "Escenario seleccionado (1=Conservador, 2=Base, 3=Optimista, 4=Estrés)")
put(A, "C4", 2, "input")
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="4", allow_blank=False)
A.add_data_validation(dv)
dv.add("C4")
put(A, "D4", '=INDEX({"Conservador","Base","Optimista","Estrés"},C4)', "out")
SEL = "Assumptions!$C$4"

G = {}  # parámetros generales -> referencia absoluta
general = [
    # key, label, unit, value, fmt, estado, fuente
    ("start_year", "Primer año del modelo", "año", 2027, "0", "HIPÓTESIS", "Calendario del promotor"),
    ("capital", "Capital aportado por el promotor (año 1)", "€", 50000, EUR, "HIPÓTESIS", "Comparar 15k–100k en 'Scenarios'"),
    ("vat", "IVA tipo general", "%", 0.21, PCT, "VERIFICADO", "Ley 37/1992"),
    ("iee_rate", "Impuesto especial cigarros: tipo ad valorem sobre PVP", "%", 0.158, PCT, "VERIFICADO", "S14"),
    ("iee_min", "Impuesto especial cigarros: mínimo", "€/1.000 u", 47, NUM2, "VERIFICADO", "S14"),
    ("retail_margin", "Margen del expendedor sobre PVP (cigarros)", "%", 0.09, PCT, "PROBABLE", "S15 (confirmar texto consolidado)"),
    ("duty", "Derecho arancelario UE 2402 10 00 (origen Cuba, sin preferencia)", "% CIF", 0.26, PCT, "PROBABLE", "S32 (confirmar TARIC)"),
    ("insurance", "Seguro de transporte", "% valor", 0.005, PCT, "HIPÓTESIS", "Cotizar con transitario"),
    ("fiscal_marks", "Marcas fiscales + identificador único trazabilidad", "€/u", 0.03, EUR2, "HIPÓTESIS", "S27, S28 (coste a confirmar con FNMT/emisor)"),
    ("customs_fee", "Despacho aduanero + agente + inspección por envío", "€/envío", 450, EUR, "HIPÓTESIS", "Cotizar agente de aduanas"),
    ("shipments", "Envíos (importaciones) por año con actividad", "nº", 2, NUM, "HIPÓTESIS", "Política de compras"),
    ("fx", "Tipo de cambio USD por EUR", "USD/€", 1.15, NUM2, "HIPÓTESIS", "Si la fábrica cotiza en EUR, poner 1,00"),
    ("units_per_box", "Unidades por caja", "u", 10, NUM, "HIPÓTESIS", "Formato propuesto"),
    ("boxes_per_pos", "Cajas vendidas por punto de venta activo y año", "cajas", 4, NUM, "HIPÓTESIS", "Benchmark interno; validar con estancos especialistas"),
    ("cover_months", "Cobertura de stock objetivo a cierre de año", "meses", 6, NUM, "HIPÓTESIS", "Maduración + plazo de reposición"),
    ("advance", "Pago anticipado a la fábrica al pedido", "%", 0.5, PCT, "HIPÓTESIS", "Habitual en contrapartes estatales cubanas; negociar"),
    ("dso", "Cobro al distribuidor", "días", 45, NUM, "HIPÓTESIS", "Negociar con distribuidor habilitado"),
    ("dpo", "Pago a proveedores de gastos generales", "días", 30, NUM, "HIPÓTESIS", ""),
    ("iee_at_import", "IEE pagado al importar (1) o diferido en depósito fiscal (0)", "1/0", 1, "0", "HIPÓTESIS", "Conservador = 1"),
    ("vat_lag", "Fracción anual de IVA pendiente a cierre (liquidación trimestral)", "fracción", 0.25, NUM2, "HIPÓTESIS", ""),
    ("royalty", "Royalty a Arnaldo por nombre/firma (% ingresos de la marca antes de IEE)", "%", 0.0, PCT, "HIPÓTESIS", "0% = opción B (sólo equity, recomendada). Opción A: 2% → VAN Optimista −4.561 € y TIR 16,4% (cálculo 10/10/2026)"),
    ("arnaldo_equity", "Participación de Arnaldo en la SL (opción B, con vesting)", "%", 0.10, PCT, "HIPÓTESIS", "Acuerdo con Arnaldo; reparto de valor, no coste"),
    ("tax_rate", "Impuesto sobre Sociedades (España)", "%", 0.25, PCT, "PROBABLE", "N04 (15% nueva creación no aplicado: conservador)"),
    ("interest_rate", "Coste de la financiación ajena (déficit de caja)", "%", 0.07, PCT, "HIPÓTESIS", "Póliza / préstamo socio"),
    ("min_cash", "Caja mínima operativa", "€", 5000, EUR, "HIPÓTESIS", ""),
    ("amort_years", "Vida útil intangibles (marca, diseño)", "años", 5, NUM, "HIPÓTESIS", ""),
    ("wacc", "Tasa de descuento (riesgo alto, no cotizada)", "%", 0.18, PCT, "HIPÓTESIS", "Prima por riesgo país/regulatorio"),
    ("exit_mult", "Múltiplo EV/EBITDA en 2031 (valor terminal)", "x", 5, X, "HIPÓTESIS", "Conservador para marca pequeña"),
]
r = 6
section(A, r, "1. Parámetros generales", 7)
header(A, r + 1, ["Parámetro", "Unidad", "Valor", "Estado", "Fuente / nota"])
r += 2
for key, lab, unit, val, fmt, st, src in general:
    put(A, f"A{r}", lab)
    put(A, f"B{r}", unit)
    put(A, f"C{r}", val, "input", fmt)
    put(A, f"D{r}", st, "flag" if st == "BLOQUEADO" else "calc")
    put(A, f"E{r}", src, "note")
    G[key] = f"Assumptions!$C${r}"
    r += 1

r += 1
section(A, r, "2. Parámetros por escenario", 9)
header(A, r + 1, ["Parámetro", "Unidad"] + SCEN + ["Seleccionado", "Estado", "Fuente / nota"])
r += 2
S = {}  # key -> list of 4 refs
scen_params = [
    ("pvp", "PVP por unidad (IVA e IEE incluidos)", "€/u", [20, 24, 28, 24], EUR2, "HIPÓTESIS", "Benchmarks S37/S38: Partagás D4 ≈21,20 €; Montecristo No.2 21,80 €; Davidoff Puro Dominicano 15,50 €. Ver rejilla 12–35 € en Unit Economics"),
    ("fob_usd", "Precio de fabricación FOB/FCA por unidad (sin cotización)", "USD/u", [7.0, 6.0, 5.0, 7.5], NUM2, "BLOQUEADO", "Requiere cotización formal del fabricante/exportador"),
    ("pack_factor", "Factor sobre coste de packaging base ('Product Cost')", "x", [1.10, 1.00, 0.90, 1.20], NUM2, "HIPÓTESIS", ""),
    ("qc", "Control de calidad en origen (inspección, catas, laboratorio)", "€/u", [0.30, 0.25, 0.20, 0.35], EUR2, "HIPÓTESIS", ""),
    ("freight", "Flete internacional + manipulación hasta almacén", "€/u", [0.70, 0.55, 0.40, 0.90], EUR2, "HIPÓTESIS", "Aéreo en lotes pequeños"),
    ("torcedor_fee", "Honorario técnico de Arnaldo (pagado vía entidad autorizada)", "€/u", [0.40, 0.35, 0.30, 0.40], EUR2, "BLOQUEADO", "Forma legal de pago sin confirmar"),
    ("dist_fee", "Tarifa del distribuidor mayorista habilitado", "% PVP", [0.06, 0.05, 0.045, 0.06], PCT, "HIPÓTESIS", "Negociable; no regulada"),
    ("merma", "Merma y rechazos post-importación", "% compras", [0.05, 0.03, 0.02, 0.08], PCT, "HIPÓTESIS", ""),
    ("muestras", "Muestras profesionales (sujetas a impuestos)", "% ventas", [0.04, 0.03, 0.025, 0.04], PCT, "HIPÓTESIS", "Sólo B2B: sin promoción al consumidor"),
    ("lead_time", "Plazo pedido→llegada a España", "meses", [6, 5, 4, 10], NUM, "HIPÓTESIS", ""),
    ("block_year", "Año a partir del cual se bloquea el suministro", "año", [9999, 9999, 9999, 2029], "0", "HIPÓTESIS", "9999 = sin bloqueo"),
    ("opex_factor", "Factor sobre gastos fijos base ('Operating Expenses')", "x", [1.15, 1.00, 1.10, 1.20], NUM2, "HIPÓTESIS", "Optimista: más esfuerzo comercial"),
]
for key, lab, unit, vals, fmt, st, src in scen_params:
    put(A, f"A{r}", lab)
    put(A, f"B{r}", unit)
    refs = []
    for s in range(4):
        col = L(3 + s)
        put(A, f"{col}{r}", vals[s], "input", fmt)
        refs.append(f"Assumptions!${col}${r}")
    put(A, f"G{r}", f"=INDEX(C{r}:F{r},{SEL})", "out", fmt)
    put(A, f"H{r}", st, "flag" if st == "BLOQUEADO" else "calc")
    put(A, f"I{r}", src, "note")
    S[key] = refs
    S[key + "_sel"] = f"Assumptions!$G${r}"
    r += 1
widths(A, {"A": 62, "B": 11, "C": 13, "D": 13, "E": 13, "F": 13, "G": 13, "H": 12, "I": 60})
A.freeze_panes = "A6"

# ---------------------------------------------------------------- Product Cost
PC = ws_["Product Cost"]
title(PC, "Coste de producto — desglose de packaging y coste en destino",
      "Packaging base por unidad (escenario Base). Los escenarios aplican el factor de 'Assumptions'. Todos los importes sin IVA.")
header(PC, 4, ["Componente (caja de 10 u)", "Unidad", "Coste por caja (€)", "Coste por unidad (€)", "Estado", "Nota"])
pack_items = [
    ("Caja de cedro español macizo con bisagra y cierre", 11.50, "HIPÓTESIS", "Cotizar en Cuba (fabricación local) y en España"),
    ("Anilla principal (papel algodón, stamping oro, relieve)", 1.20, "HIPÓTESIS", "0,12 €/u en tiradas ≥10.000"),
    ("Pie de anilla / numeración de pliego", 0.40, "HIPÓTESIS", ""),
    ("Hoja de cedro interior + papel de seda", 0.60, "HIPÓTESIS", ""),
    ("Colofón impreso numerado (tarjeta interior)", 0.80, "HIPÓTESIS", "Elemento de trazabilidad de lote"),
    ("Etiquetado de advertencias sanitarias (laterales/tapa)", 0.70, "HIPÓTESIS", "Diseño conforme RD 579/2017"),
    ("Precinto de garantía de la marca (no fiscal)", 0.50, "HIPÓTESIS", "No imita sellos de Habanos"),
    ("Retractilado + cartón exterior de expedición (prorrata)", 0.80, "HIPÓTESIS", ""),
]
r = 5
for lab, cost, st, note in pack_items:
    put(PC, f"A{r}", lab)
    put(PC, f"B{r}", "caja")
    put(PC, f"C{r}", cost, "input", EUR2)
    put(PC, f"D{r}", f"=C{r}/{G['units_per_box']}", "calc", EUR2)
    put(PC, f"E{r}", st)
    put(PC, f"F{r}", note, "note")
    r += 1
put(PC, f"A{r}", "TOTAL packaging base")
put(PC, f"C{r}", f"=SUM(C5:C{r-1})", "out", EUR2)
put(PC, f"D{r}", f"=SUM(D5:D{r-1})", "out", EUR2)
PACK_TOTAL = f"{q('Product Cost')}!$D${r}"
r += 2
section(PC, r, "Coste en destino por unidad — escenario seleccionado (lee del motor)", 6)
r += 1
PC_START = r
widths(PC, {"A": 58, "B": 10, "C": 18, "D": 18, "E": 12, "F": 46})

# ---------------------------------------------------------------- Sales Forecast (inputs de demanda)
SF = ws_["Sales Forecast"]
title(SF, "Previsión de demanda (unidades) por escenario",
      "Demanda B2B a través del canal legal (distribuidor habilitado → expendedurías). 2027 = año de habilitaciones, sin ventas.")
header(SF, 4, ["Escenario"] + [str(y) for y in YEARS] + ["Total"])
demand = {0: [0, 1000, 1500, 2000, 2500], 1: [0, 2000, 3500, 5000, 6500],
          2: [0, 3500, 6000, 9000, 12000], 3: [0, 2000, 3500, 5000, 6500]}
DEM = {}
for s in range(4):
    rr = 5 + s
    put(SF, f"A{rr}", SCEN[s] + (" (demanda; suministro bloqueado desde 2029)" if s == 3 else ""))
    for i in range(NY):
        put(SF, f"{L(2+i)}{rr}", demand[s][i], "input", NUM)
    put(SF, f"G{rr}", f"=SUM(B{rr}:F{rr})", "calc", NUM)
    DEM[s] = [f"{q('Sales Forecast')}!${L(2+i)}${rr}" for i in range(NY)]
widths(SF, {"A": 52, "B": 13, "C": 13, "D": 13, "E": 13, "F": 13, "G": 14})

# ---------------------------------------------------------------- Operating Expenses
OX = ws_["Operating Expenses"]
title(OX, "Gastos operativos fijos y CAPEX (escenario Base, € sin IVA)",
      "El motor multiplica el total por el factor de escenario. 'Remuneración del promotor' = 0: el modelo muestra el resultado ANTES de retribuir al gestor.")
header(OX, 4, ["Partida", "Estado"] + [str(y) for y in YEARS])
opex = [
    ("Due diligence jurídica Cuba (Fase 0) — abogado cubano + español", "HIPÓTESIS", [6000, 0, 0, 0, 0]),
    ("Asesoría jurídica España: contratos, sociedad, propiedad intelectual", "HIPÓTESIS", [6000, 2500, 2500, 2500, 2500]),
    ("Habilitación CMT, notificación EU-CEG, gestión regulatoria", "HIPÓTESIS", [3000, 800, 800, 1500, 800]),
    ("Gestoría, contabilidad, auditoría ligera", "HIPÓTESIS", [1800, 3600, 3600, 3800, 4000]),
    ("Seguros (RC producto, mercancía en almacén)", "HIPÓTESIS", [600, 1500, 1500, 1600, 1700]),
    ("Almacén climatizado / depósito (3PL)", "HIPÓTESIS", [0, 1800, 2400, 3000, 3600]),
    ("Viajes a Cuba (negociación, control de calidad)", "HIPÓTESIS", [5000, 3000, 3000, 3000, 3000]),
    ("Desarrollo comercial B2B (ferias profesionales, logística de muestras)", "HIPÓTESIS", [1000, 6000, 5000, 5000, 5000]),
    ("Software, trazabilidad, registros", "HIPÓTESIS", [500, 1200, 1200, 1200, 1200]),
    ("Renovaciones y vigilancia de marcas", "HIPÓTESIS", [0, 300, 300, 300, 300]),
    ("Remuneración del promotor (no incluida)", "HIPÓTESIS", [0, 0, 0, 0, 0]),
]
r = 5
for lab, st, vals in opex:
    put(OX, f"A{r}", lab)
    put(OX, f"B{r}", st)
    for i in range(NY):
        put(OX, f"{L(3+i)}{r}", vals[i], "input", EUR)
    r += 1
put(OX, f"A{r}", "Imprevistos (% sobre el subtotal)")
put(OX, f"B{r}", 0.10, "input", PCT)
for i in range(NY):
    c = L(3 + i)
    put(OX, f"{c}{r}", f"=SUM({c}5:{c}{r-1})*$B${r}", "calc", EUR)
cont_row = r
r += 1
put(OX, f"A{r}", "TOTAL GASTOS FIJOS (base)")
for i in range(NY):
    c = L(3 + i)
    put(OX, f"{c}{r}", f"=SUM({c}5:{c}{r-1})", "out", EUR)
OPEX_TOTAL = [f"{q('Operating Expenses')}!${L(3+i)}${r}" for i in range(NY)]
r += 2
section(OX, r, "CAPEX intangible (amortizable)", 7)
r += 1
capex = [("Registro de marcas: EUIPO clase 34 + Cuba (OCPI vía agente) + vigilancia", "PROBABLE", [4000, 0, 0, 0, 0]),
         ("Diseño de identidad, artes finales, troqueles y planchas", "HIPÓTESIS", [9000, 0, 3000, 0, 0])]
c0 = r
for lab, st, vals in capex:
    put(OX, f"A{r}", lab)
    put(OX, f"B{r}", st)
    for i in range(NY):
        put(OX, f"{L(3+i)}{r}", vals[i], "input", EUR)
    r += 1
put(OX, f"A{r}", "TOTAL CAPEX")
for i in range(NY):
    c = L(3 + i)
    put(OX, f"{c}{r}", f"=SUM({c}{c0}:{c}{r-1})", "out", EUR)
CAPEX_TOTAL = [f"{q('Operating Expenses')}!${L(3+i)}${r}" for i in range(NY)]
widths(OX, {"A": 66, "B": 12, "C": 13, "D": 13, "E": 13, "F": 13, "G": 13})

# ---------------------------------------------------------------- Engine
E = ws_["Engine"]
title(E, "Motor de cálculo — 4 escenarios en paralelo (no editar: sólo fórmulas)",
      "Bloques: Conservador C:G · Base I:M · Optimista O:S · Estrés U:Y. Las hojas de presentación leen el escenario seleccionado.")
for s in range(4):
    c0 = 3 + s * 6
    E.merge_cells(start_row=3, start_column=c0, end_row=3, end_column=c0 + 4)
    cc = E.cell(row=3, column=c0, value=SCEN[s])
    cc.font, cc.fill, cc.alignment = F_H, FILL_H, CENTER
    for i in range(NY):
        hc = E.cell(row=4, column=c0 + i, value=YEARS[i])
        hc.font, hc.fill, hc.alignment = F_OUT, FILL_SEC, CENTER
E["A4"], E["B4"] = "Concepto", "Unidad"

rows = [
    ("_sec", "VOLUMEN E INVENTARIO (unidades)"),
    ("year", "Año", "0", "{g:start_year}+{i}"),
    ("demand", "Demanda", NUM, "{dem}"),
    ("supply_ok", "Suministro disponible (1/0)", "0", "IF({r:year}<{s:block_year},1,0)"),
    ("k", "Factor de consumo por unidad vendida (merma+muestras)", "0.000", "1/(1-{s:merma})+{s:muestras}"),
    ("inv_begin", "Stock inicial", NUM, "{prev:inv_end}"),
    ("target_end", "Stock objetivo a cierre", NUM, "{g:cover_months}/12*{dem_next}*{r:supply_ok}"),
    ("purchased", "Unidades compradas (importadas)", NUM, "IF({r:supply_ok}=1,MAX(0,{r:demand}*{r:k}+{r:target_end}-{r:inv_begin}),0)"),
    ("sold", "Unidades vendidas", NUM, "IF({r:supply_ok}=1,{r:demand},MIN({r:demand},{r:inv_begin}/{r:k}))"),
    ("consumed", "Unidades consumidas (ventas+merma+muestras)", NUM, "{r:sold}*{r:k}"),
    ("inv_end", "Stock final", NUM, "{r:inv_begin}+{r:purchased}-{r:consumed}"),
    ("_sec", "ECONOMÍA UNITARIA (€/u)"),
    ("pvp", "PVP", EUR2, "{s:pvp}"),
    ("fob_eur", "Fabricación FOB en €", EUR2, "{s:fob_usd}/{g:fx}"),
    ("pack", "Packaging", EUR2, PACK_TOTAL + "*{s:pack_factor}"),
    ("cif_u", "Valor CIF (base aduanera)", EUR2, "({r:fob_eur}+{r:pack}+{s:qc}+{s:freight})*(1+{g:insurance})"),
    ("duty_u", "Arancel", EUR2, "{r:cif_u}*{g:duty}"),
    ("landed_u", "Coste en destino (sin IEE ni IVA)", EUR2, "{r:cif_u}+{r:duty_u}+{g:fiscal_marks}+{s:torcedor_fee}"),
    ("iee_u", "Impuesto especial por unidad", EUR2, "MAX({g:iee_rate}*{r:pvp},{g:iee_min}/1000)"),
    ("unit_contrib", "Contribución por unidad vendida", EUR2,
     "({r:pvp}/(1+{g:vat})-{r:pvp}*({g:retail_margin}+{s:dist_fee}))*(1-{g:royalty})-{r:k}*({r:iee_u}+{r:landed_u})"),
    ("_sec", "CUENTA DE RESULTADOS (€)"),
    ("rev_pvp", "Ventas a PVP (consumidor, IVA incl.) — informativo", EUR, "{r:sold}*{r:pvp}"),
    ("rev_gross", "Ventas a PVP sin IVA", EUR, "{r:sold}*{r:pvp}/(1+{g:vat})"),
    ("retail_margin", "(-) Margen expendedor", EUR, "{r:sold}*{r:pvp}*{g:retail_margin}"),
    ("dist_fee", "(-) Tarifa distribuidor", EUR, "{r:sold}*{r:pvp}*{s:dist_fee}"),
    ("net_rev", "Ingresos de la marca (antes de IEE)", EUR, "{r:rev_gross}-{r:retail_margin}-{r:dist_fee}"),
    ("iee_exp", "(-) Impuesto especial devengado", EUR, "{r:consumed}*{r:iee_u}"),
    ("net_rev_iee", "Ingresos netos de impuestos", EUR, "{r:net_rev}-{r:iee_exp}"),
    ("cogs", "(-) Coste de ventas (incl. merma y muestras)", EUR, "{r:consumed}*{r:landed_u}"),
    ("customs_fees", "(-) Gastos de despacho aduanero", EUR, "IF({r:purchased}>0,{g:shipments}*{g:customs_fee},0)"),
    ("royalty", "(-) Royalty de nombre a Arnaldo", EUR, "{r:net_rev}*{g:royalty}"),
    ("gross_margin", "Margen bruto", EUR, "{r:net_rev_iee}-{r:cogs}-{r:customs_fees}-{r:royalty}"),
    ("opex", "(-) Gastos fijos", EUR, "{opex}*{s:opex_factor}"),
    ("ebitda", "EBITDA", EUR, "{r:gross_margin}-{r:opex}"),
    ("capex", "CAPEX intangible", EUR, "{capex}"),
    ("cum_capex", "CAPEX acumulado", EUR, "{prev:cum_capex}+{r:capex}"),
    ("da", "(-) Amortización", EUR, "{r:cum_capex}/{g:amort_years}"),
    ("cum_da", "Amortización acumulada", EUR, "{prev:cum_da}+{r:da}"),
    ("intang_net", "Intangible neto", EUR, "{r:cum_capex}-{r:cum_da}"),
    ("ebit", "EBIT", EUR, "{r:ebitda}-{r:da}"),
    ("interest", "(-) Gastos financieros (sobre deuda inicial del año)", EUR, "{prev:debt}*{g:interest_rate}"),
    ("ebt", "Resultado antes de impuestos", EUR, "{r:ebit}-{r:interest}"),
    ("bin_begin", "Bases negativas pendientes (inicio)", EUR, "{prev:bin_end}"),
    ("taxable", "Base imponible tras compensación", EUR, "MAX(0,{r:ebt}-{r:bin_begin})"),
    ("bin_end", "Bases negativas pendientes (cierre)", EUR, "MAX(0,{r:bin_begin}-MAX({r:ebt},0))+MAX(-{r:ebt},0)"),
    ("tax", "(-) Impuesto sobre Sociedades", EUR, "{r:taxable}*{g:tax_rate}"),
    ("net_income", "Resultado neto", EUR, "{r:ebt}-{r:tax}"),
    ("_sec", "CAPITAL CIRCULANTE (€, saldos a cierre)"),
    ("inv_value", "Existencias (coste en destino)", EUR, "{r:inv_end}*{r:landed_u}"),
    ("iee_paid", "IEE pagado en el año", EUR, "IF({g:iee_at_import}=1,{r:purchased},{r:consumed})*{r:iee_u}"),
    ("iee_prepaid", "IEE anticipado sobre existencias", EUR, "IF({g:iee_at_import}=1,{r:inv_end}*{r:iee_u},0)"),
    ("ar", "Clientes (distribuidor)", EUR, "{r:net_rev}*(1+{g:vat})*{g:dso}/365"),
    ("prepay", "Anticipos a fábrica (pedidos en curso)", EUR, "{g:advance}*{next:purchased}*({next:fob_eur}+{next:pack}+{s:qc})*MIN(1,{s:lead_time}/12)"),
    ("vat_import", "IVA de importación soportado en el año", EUR, "({r:purchased}*({r:cif_u}+{r:duty_u})+{r:iee_paid})*{g:vat}"),
    ("vat_rec", "Hacienda deudora por IVA", EUR, "{r:vat_import}*{g:vat_lag}"),
    ("ap", "Proveedores de gastos generales", EUR, "{r:opex}*{g:dpo}/365"),
    ("vat_out", "IVA repercutido en el año", EUR, "{r:net_rev}*{g:vat}"),
    ("vat_pay", "Hacienda acreedora por IVA", EUR, "{r:vat_out}*{g:vat_lag}"),
    ("nwc", "Capital circulante neto", EUR, "{r:inv_value}+{r:iee_prepaid}+{r:ar}+{r:prepay}+{r:vat_rec}-{r:ap}-{r:vat_pay}"),
    ("d_nwc", "Variación del circulante", EUR, "{r:nwc}-{prev:nwc}"),
    ("_sec", "FLUJOS DE CAJA Y FINANCIACIÓN (€)"),
    ("cfo", "Flujo de explotación", EUR, "{r:net_income}+{r:da}-{r:d_nwc}"),
    ("cfi", "Flujo de inversión", EUR, "-{r:capex}"),
    ("equity_in", "Aportación de capital", EUR, "IF({i}=0,{g:capital},0)"),
    ("ncb", "Posición neta de caja antes de financiación ajena", EUR, "{prev:cash}-{prev:debt}+{r:cfo}+{r:cfi}+{r:equity_in}"),
    ("debt", "Financiación ajena necesaria (saldo)", EUR, "MAX(0,{g:min_cash}-{r:ncb})"),
    ("d_debt", "Disposición / (devolución) de financiación", EUR, "{r:debt}-{prev:debt}"),
    ("cash", "Caja final", EUR, "{r:ncb}+{r:debt}"),
    ("_sec", "BALANCE (€)"),
    ("total_assets", "Total activo", EUR, "{r:cash}+{r:inv_value}+{r:iee_prepaid}+{r:ar}+{r:prepay}+{r:vat_rec}+{r:intang_net}"),
    ("total_liab", "Total pasivo exigible", EUR, "{r:ap}+{r:vat_pay}+{r:debt}"),
    ("cum_ni", "Resultados acumulados", EUR, "{prev:cum_ni}+{r:net_income}"),
    ("equity", "Patrimonio neto", EUR, "{g:capital}+{r:cum_ni}"),
    ("check", "CONTROL: Activo − Pasivo − PN (debe ser 0)", NUM2, "ROUND({r:total_assets}-{r:total_liab}-{r:equity},2)"),
    ("_sec", "RENTABILIDAD (€)"),
    ("fcf", "Flujo de caja libre del proyecto (sin financiación)", EUR, "{r:ebitda}-{r:tax}-{r:d_nwc}-{r:capex}"),
    ("cum_fcf", "FCF acumulado", EUR, "{prev:cum_fcf}+{r:fcf}"),
    ("tv", "Valor terminal (sólo 2031; EBITDA>0 × múltiplo)", EUR, "IF({i}=4,MAX(0,{r:ebitda})*{g:exit_mult},0)"),
    ("fcf_tv", "FCF + valor terminal", EUR, "{r:fcf}+{r:tv}"),
    ("df", "Factor de descuento (fin de año)", "0.0000", "1/(1+{g:wacc})^({i}+1)"),
    ("pv", "Valor actual", EUR, "{r:fcf_tv}*{r:df}"),
    ("payback_flag", "Año de recuperación (si ocurre)", "0", "IF(AND({r:cum_fcf}>=0,{prevneg:cum_fcf}<0),{r:year},\"\")"),
    ("nopat", "NOPAT (EBIT × (1−t))", EUR, "{r:ebit}*(1-{g:tax_rate})"),
    ("invested", "Capital invertido (circulante + intangible)", EUR, "{r:nwc}+{r:intang_net}"),
    ("roic", "ROIC (NOPAT / capital invertido medio)", PCT, "IF(({r:invested}+{prev:invested})/2>1000,{r:nopat}/(({r:invested}+{prev:invested})/2),\"n/a\")"),
    ("be_ebitda", "Punto muerto EBITDA (unidades/año)", NUM, "IF({r:unit_contrib}>0,({r:opex}+{g:shipments}*{g:customs_fee})/{r:unit_contrib},\"n/a\")"),
    ("be_ebit", "Punto muerto EBIT (unidades/año)", NUM, "IF({r:unit_contrib}>0,({r:opex}+{g:shipments}*{g:customs_fee}+{r:da})/{r:unit_contrib},\"n/a\")"),
    ("pos_needed", "Puntos de venta activos necesarios", NUM, "{r:sold}/({g:boxes_per_pos}*{g:units_per_box})"),
]
ROW = {}
r = 5
for spec in rows:
    if spec[0] == "_sec":
        r += 1
        section(E, r, spec[1], 25)
        r += 1
        continue
    ROW[spec[0]] = r
    r += 1


def col(s, i):
    return L(3 + s * 6 + i)


import re


def resolve(t, s, i):
    def rep(m):
        kind, key = m.group(1), m.group(2)
        if kind == "r":
            return f"{col(s, i)}{ROW[key]}"
        if kind == "prev":
            return "0" if i == 0 else f"{col(s, i-1)}{ROW[key]}"
        if kind == "prevneg":
            return "-1" if i == 0 else f"{col(s, i-1)}{ROW[key]}"
        if kind == "next":
            return f"{col(s, min(i+1, NY-1))}{ROW[key]}"
        if kind == "g":
            return G[key]
        if kind == "s":
            return S[key][s]
        raise KeyError(kind)
    t = re.sub(r"\{(r|prev|prevneg|next|g|s):([a-z_0-9]+)\}", rep, t)
    t = t.replace("{dem}", DEM[s][i]).replace("{dem_next}", DEM[s][min(i + 1, NY - 1)])
    t = t.replace("{opex}", OPEX_TOTAL[i]).replace("{capex}", CAPEX_TOTAL[i]).replace("{i}", str(i))
    return "=" + t


for spec in rows:
    if spec[0] == "_sec":
        continue
    key, lab, fmt, tmpl = spec
    rr = ROW[key]
    E.cell(row=rr, column=1, value=lab).font = F_OUT if key in ("ebitda", "net_income", "cash", "check", "fcf", "nwc") else F_CALC
    for s in range(4):
        for i in range(NY):
            c = E.cell(row=rr, column=3 + s * 6 + i, value=resolve(tmpl, s, i))
            c.number_format = fmt
            c.font = F_CALC
E.column_dimensions["A"].width = 58
E.column_dimensions["B"].width = 4
for s in range(4):
    for i in range(NY):
        E.column_dimensions[col(s, i)].width = 12
    E.column_dimensions[L(3 + s * 6 + 5)].width = 3
E.freeze_panes = "C5"
E.conditional_formatting.add(f"C{ROW['check']}:Y{ROW['check']}", CellIsRule(operator="notEqual", formula=["0"], fill=FILL_FLAG))


def SELREF(key, i):
    rr = ROW[key]
    return f"INDEX(Engine!$C${rr}:$Y${rr},1,({SEL}-1)*6+{i+1})"


def ENG(key, s, i):
    return f"Engine!${col(s, i)}${ROW[key]}"


def ENGROW(key, s):
    return f"Engine!${col(s,0)}${ROW[key]}:${col(s,NY-1)}${ROW[key]}"


def statement(ws, start, items, ytitle=True):
    """items: (label, key or None, fmt, kind)"""
    if ytitle:
        header(ws, start, ["Concepto"] + [str(y) for y in YEARS])
        start += 1
    for lab, key, fmt, kind in items:
        if key is None:
            section(ws, start, lab, 6)
        else:
            ws.cell(row=start, column=1, value=lab).font = F_OUT if kind == "out" else F_CALC
            for i in range(NY):
                put(ws, f"{L(2+i)}{start}", "=" + SELREF(key, i), kind, fmt)
        start += 1
    return start


# Product Cost: tabla en destino (seleccionado)
pc_items = [("Fabricación FOB (€)", "fob_eur"), ("Packaging", "pack"), ("Valor CIF (incl. QC, flete, seguro)", "cif_u"),
            ("Arancel", "duty_u"), ("Coste en destino (incl. marcas fiscales y honorario técnico)", "landed_u"),
            ("Impuesto especial por unidad", "iee_u"), ("Contribución por unidad vendida", "unit_contrib")]
r = PC_START
header(PC, r, ["Concepto (€/u)"] + [str(y) for y in YEARS])
r += 1
for lab, key in pc_items:
    PC.cell(row=r, column=1, value=lab)
    for i in range(NY):
        put(PC, f"{L(2+i)}{r}", "=" + SELREF(key, i), "out" if key in ("landed_u", "unit_contrib") else "calc", EUR2)
    r += 1
put(PC, f"A{r+1}", "Coste en destino por caja de 10 (2028)")
put(PC, f"C{r+1}", f"=C{r-3}*{G['units_per_box']}", "out", EUR2)

# ---------------------------------------------------------------- Import & Taxes
IT = ws_["Import & Taxes"]
title(IT, "Importación y fiscalidad — simulación de un envío (DUA) y reparto del PVP",
      "Escenario seleccionado. El IVA de importación es recuperable (impacto de caja, no de coste). Base IVA importación = valor en aduana + arancel + IEE.")
put(IT, "A4", "Unidades del envío")
put(IT, "C4", 3000, "input", NUM)
items = [
    ("Valor en aduana (CIF) por unidad", f"={SELREF('cif_u',1)}", EUR2),
    ("Valor en aduana del envío", "=C4*C5", EUR),
    ("Derecho arancelario", f"=C6*{G['duty']}", EUR),
    ("Impuesto especial (si se paga al importar)", f"=C4*{SELREF('iee_u',1)}*{G['iee_at_import']}", EUR),
    ("Base imponible IVA importación", "=C6+C7+C8", EUR),
    ("IVA de importación (recuperable)", f"=C9*{G['vat']}", EUR),
    ("Despacho y agente de aduanas", f"={G['customs_fee']}", EUR),
    ("PAGO TOTAL EN FRONTERA", "=C7+C8+C10+C11", EUR),
    ("  de ello, coste definitivo (arancel + despacho)", "=C7+C11", EUR),
    ("  de ello, impuesto recuperable vía precio de venta (IEE)", "=C8", EUR),
    ("  de ello, IVA recuperable vía liquidación (≈1 trimestre)", "=C10", EUR),
]
r = 5
for lab, f, fmt in items:
    put(IT, f"A{r}", lab)
    put(IT, f"C{r}", f, "out" if "TOTAL" in lab else "calc", fmt)
    r += 1
r += 1
section(IT, r, "Reparto del PVP de una unidad (año 2028, escenario seleccionado)", 6)
r += 1
header(IT, r, ["Destino del PVP", "", "€/u", "% PVP"])
r += 1
p0 = r
split = [("IVA (Hacienda)", f"={SELREF('pvp',1)}*{G['vat']}/(1+{G['vat']})"),
         ("Impuesto especial (Hacienda)", f"={SELREF('iee_u',1)}"),
         ("Margen expendedor", f"={SELREF('pvp',1)}*{G['retail_margin']}"),
         ("Tarifa distribuidor", f"={SELREF('pvp',1)}*{S['dist_fee_sel']}"),
         ("Arancel (Hacienda UE)", f"={SELREF('duty_u',1)}"),
         ("Fabricación, packaging, QC, flete, seguro", f"={SELREF('cif_u',1)}"),
         ("Marcas fiscales + honorario técnico", f"={SELREF('landed_u',1)}-{SELREF('cif_u',1)}-{SELREF('duty_u',1)}"),
         ("Merma y muestras (coste + IEE)", f"=({SELREF('k',1)}-1)*({SELREF('iee_u',1)}+{SELREF('landed_u',1)})"),
         ("Royalty de nombre a Arnaldo", f"=({SELREF('pvp',1)}/(1+{G['vat']})-{SELREF('pvp',1)}*({G['retail_margin']}+{S['dist_fee_sel']}))*{G['royalty']}"),
         ("Contribución de la marca (antes de gastos fijos)", f"={SELREF('unit_contrib',1)}")]
for lab, f in split:
    put(IT, f"A{r}", lab)
    put(IT, f"C{r}", f, "out" if "Contribución" in lab else "calc", EUR2)
    put(IT, f"D{r}", f"=C{r}/{SELREF('pvp',1)}", "calc", PCT)
    r += 1
put(IT, f"A{r}", "Control: suma = PVP")
put(IT, f"C{r}", f"=SUM(C{p0}:C{r-1})", "out", EUR2)
put(IT, f"D{r}", f"=SUM(D{p0}:D{r-1})", "out", PCT)
widths(IT, {"A": 60, "B": 4, "C": 16, "D": 12})

# ---------------------------------------------------------------- Unit Economics (rejilla PVP)
UE = ws_["Unit Economics"]
title(UE, "Economía unitaria por nivel de PVP (costes del escenario seleccionado, año 2028)",
      "Cada columna recalcula impuestos y márgenes para el PVP indicado. Las cifras del distribuidor, expendedor y marca se separan explícitamente.")
prices = [12, 15, 18, 20, 25, 30, 35]
header(UE, 4, ["Concepto (€/u)"] + [f"PVP {p} €" for p in prices])
for j, p in enumerate(prices):
    put(UE, f"{L(2+j)}5", p, "input", EUR2)
UE["A5"] = "PVP (input)"
k = SELREF("k", 1)
lines = [
    ("IVA incluido en PVP", "={c}5*{vat}/(1+{vat})"),
    ("Impuesto especial (15,8% PVP, mín. 0,047 €)", "=MAX({iee}*{c}5,{iemin}/1000)"),
    ("Margen expendedor", "={c}5*{rm}"),
    ("Tarifa distribuidor", "={c}5*{df}"),
    ("Ingreso de la marca antes de IEE", "={c}5/(1+{vat})-{c}8-{c}9"),
    ("Ingreso de la marca neto de IEE", "={c}10-{c}7"),
    ("Coste en destino (sin IEE)", "=" + SELREF("landed_u", 1)),
    ("Merma y muestras (coste + IEE)", "=(" + k + "-1)*({c}7+{c}12)"),
    ("CONTRIBUCIÓN por unidad (tras royalty Arnaldo)", "={c}11-{c}12-{c}13-{c}10*" + G["royalty"]),
    ("Margen de contribución sobre ingreso neto de IEE", "=IF({c}11<>0,{c}14/{c}11,0)"),
    ("Contribución sobre PVP", "={c}14/{c}5"),
    ("Punto muerto EBITDA 2030 (unidades/año)", "=IF({c}14>0,(" + SELREF("opex", 3) + "+" + G["shipments"] + "*" + G["customs_fee"] + ")/{c}14,\"no viable\")"),
    ("FOB máximo (USD) para contribución ≥ 25% del ingreso neto", "=MAX(0,(({c}11*0.75-{c}10*" + G["royalty"] + "-{c}7*(" + k + "-1))/" + k + "-(" + SELREF("landed_u", 1) + "-" + SELREF("fob_eur", 1) + "*(1+" + G["insurance"] + ")*(1+" + G["duty"] + ")))/((1+" + G["insurance"] + ")*(1+" + G["duty"] + "))*" + G["fx"] + ")"),
]
for n, (lab, f) in enumerate(lines):
    rr = 6 + n
    UE.cell(row=rr, column=1, value=lab).font = F_OUT if "CONTRIB" in lab else F_CALC
    for j in range(len(prices)):
        c = L(2 + j)
        ff = f.format(c=c, vat=G["vat"], iee=G["iee_rate"], iemin=G["iee_min"], rm=G["retail_margin"], df=S["dist_fee_sel"])
        fmt = PCT if ("Margen de contribución" in lab or "sobre PVP" in lab) else (NUM if "Punto muerto" in lab else EUR2)
        put(UE, f"{c}{rr}", ff, "out" if "CONTRIB" in lab or "Punto" in lab else "calc", fmt)
widths(UE, {"A": 60, "B": 12, "C": 12, "D": 12, "E": 12, "F": 12, "G": 12, "H": 12})
put(UE, "A21", "Lectura: el Estado (IVA+IEE+arancel) captura ~35-40% del PVP; el canal (expendedor+distribuidor) ~14%. La marca debe pagar fábrica, logística y estructura con el resto.", "note")

# Sales forecast derivado (seleccionado)
r = 11
section(SF, r, "Escenario seleccionado — volumen derivado y esfuerzo comercial", 7)
r = statement(SF, r + 1, [("Unidades vendidas", "sold", NUM, "out"), ("Unidades importadas", "purchased", NUM, "calc"),
                          ("Stock final (unidades)", "inv_end", NUM, "calc"), ("Cajas vendidas", None, NUM, "calc"),
                          ("Puntos de venta activos necesarios", "pos_needed", NUM, "out"),
                          ("Ventas a PVP (consumidor)", "rev_pvp", EUR, "calc")])
# reemplazar la fila 'Cajas vendidas' (sección) por fórmula
cajas_row = r - 3
for i in range(NY):
    put(SF, f"{L(2+i)}{cajas_row}", f"={L(2+i)}{cajas_row-3}/{G['units_per_box']}", "calc", NUM)
SF.cell(row=cajas_row, column=1, value="Cajas vendidas").font = F_CALC
for cc in range(1, 7):
    SF.cell(row=cajas_row, column=cc).fill = PatternFill(fill_type=None)

# ---------------------------------------------------------------- P&L
PL = ws_["Profit & Loss"]
title(PL, "Cuenta de resultados 2027–2031 (escenario seleccionado)", "€ sin IVA. Los importes del consumidor y del canal se muestran para separar ingresos de marca, distribuidor y detallista.")
put(PL, "A3", "=\"Escenario: \"&Assumptions!$D$4", "out")
statement(PL, 5, [
    ("Unidades vendidas", "sold", NUM, "calc"),
    ("Ventas a PVP con IVA (gasto del consumidor, informativo)", "rev_pvp", EUR, "calc"),
    ("Ventas a PVP sin IVA", "rev_gross", EUR, "calc"),
    ("(-) Margen del expendedor (ingreso del detallista)", "retail_margin", EUR, "calc"),
    ("(-) Tarifa del distribuidor (ingreso del mayorista)", "dist_fee", EUR, "calc"),
    ("INGRESOS DE LA MARCA (antes de IEE)", "net_rev", EUR, "out"),
    ("(-) Impuesto especial", "iee_exp", EUR, "calc"),
    ("Ingresos netos de impuestos", "net_rev_iee", EUR, "out"),
    ("(-) Coste de ventas", "cogs", EUR, "calc"),
    ("(-) Despachos aduaneros", "customs_fees", EUR, "calc"),
    ("(-) Royalty de nombre a Arnaldo", "royalty", EUR, "calc"),
    ("MARGEN BRUTO", "gross_margin", EUR, "out"),
    ("(-) Gastos fijos", "opex", EUR, "calc"),
    ("EBITDA", "ebitda", EUR, "out"),
    ("(-) Amortización", "da", EUR, "calc"),
    ("EBIT", "ebit", EUR, "out"),
    ("(-) Gastos financieros", "interest", EUR, "calc"),
    ("Resultado antes de impuestos", "ebt", EUR, "calc"),
    ("(-) Impuesto sobre Sociedades", "tax", EUR, "calc"),
    ("RESULTADO NETO", "net_income", EUR, "out"),
])
r = 26
put(PL, f"A{r}", "Margen bruto / ingresos netos de impuestos")
for i in range(NY):
    c = L(2 + i)
    put(PL, f"{c}{r}", f"=IF({c}13<>0,{c}17/{c}13,0)", "calc", PCT)
put(PL, f"A{r+1}", "Margen EBITDA / ingresos de la marca")
for i in range(NY):
    c = L(2 + i)
    put(PL, f"{c}{r+1}", f"=IF({c}11<>0,{c}19/{c}11,0)", "calc", PCT)
widths(PL, {"A": 62, "B": 14, "C": 14, "D": 14, "E": 14, "F": 14})

# ---------------------------------------------------------------- Cash Flow
CF = ws_["Cash Flow"]
title(CF, "Estado de flujos de caja 2027–2031 (escenario seleccionado)", "Método indirecto. La financiación ajena se dimensiona automáticamente para mantener la caja mínima (no circular).")
put(CF, "A3", "=\"Escenario: \"&Assumptions!$D$4", "out")
statement(CF, 5, [
    ("Resultado neto", "net_income", EUR, "calc"),
    ("(+) Amortización", "da", EUR, "calc"),
    ("Variación del circulante (aumento = consumo de caja; se resta)", "d_nwc", EUR, "calc"),
    ("FLUJO DE EXPLOTACIÓN", "cfo", EUR, "out"),
    ("FLUJO DE INVERSIÓN", "cfi", EUR, "out"),
    ("Aportación de capital", "equity_in", EUR, "calc"),
    ("Disposición / (devolución) de financiación ajena", "d_debt", EUR, "calc"),
    ("VARIACIÓN NETA DE CAJA", None, EUR, "out"),
    ("Caja final", "cash", EUR, "out"),
    ("Saldo de financiación ajena necesaria", "debt", EUR, "out"),
    ("Flujo de caja libre del proyecto", "fcf", EUR, "out"),
    ("FCF acumulado", "cum_fcf", EUR, "calc"),
])
# variación neta de caja (fila 13)
CF.cell(row=13, column=1, value="VARIACIÓN NETA DE CAJA").font = F_OUT
for cc in range(1, 7):
    CF.cell(row=13, column=cc).fill = PatternFill(fill_type=None)
for i in range(NY):
    c = L(2 + i)
    put(CF, f"{c}13", f"=SUM({c}9:{c}12)", "out", EUR)
put(CF, "A19", "Control: caja final − caja inicial − variación neta")
for i in range(NY):
    c = L(2 + i)
    prev = "0" if i == 0 else f"{L(1+i)}14"
    put(CF, f"{c}19", f"=ROUND({c}14-{prev}-{c}13,2)", "calc", NUM2)
widths(CF, {"A": 58, "B": 14, "C": 14, "D": 14, "E": 14, "F": 14})

# ---------------------------------------------------------------- Balance Sheet
BS = ws_["Balance Sheet"]
title(BS, "Balance simplificado a 31/12 (escenario seleccionado)", "")
put(BS, "A3", "=\"Escenario: \"&Assumptions!$D$4", "out")
statement(BS, 5, [
    ("ACTIVO", None, EUR, "calc"),
    ("Intangible neto (marcas, diseño)", "intang_net", EUR, "calc"),
    ("Existencias", "inv_value", EUR, "calc"),
    ("IEE anticipado sobre existencias", "iee_prepaid", EUR, "calc"),
    ("Anticipos a fábrica", "prepay", EUR, "calc"),
    ("Clientes", "ar", EUR, "calc"),
    ("Hacienda deudora IVA", "vat_rec", EUR, "calc"),
    ("Tesorería", "cash", EUR, "calc"),
    ("TOTAL ACTIVO", "total_assets", EUR, "out"),
    ("PATRIMONIO NETO Y PASIVO", None, EUR, "calc"),
    ("Capital", None, EUR, "calc"),
    ("Resultados acumulados", "cum_ni", EUR, "calc"),
    ("PATRIMONIO NETO", "equity", EUR, "out"),
    ("Financiación ajena", "debt", EUR, "calc"),
    ("Proveedores", "ap", EUR, "calc"),
    ("Hacienda acreedora IVA", "vat_pay", EUR, "calc"),
    ("TOTAL PN + PASIVO", None, EUR, "out"),
    ("CONTROL (debe ser 0)", "check", NUM2, "out"),
])
BS.cell(row=16, column=1, value="Capital").font = F_CALC
for cc in range(1, 7):
    BS.cell(row=16, column=cc).fill = PatternFill(fill_type=None)
    BS.cell(row=22, column=cc).fill = PatternFill(fill_type=None)
BS.cell(row=22, column=1, value="TOTAL PN + PASIVO").font = F_OUT
for i in range(NY):
    c = L(2 + i)
    put(BS, f"{c}16", f"={G['capital']}", "calc", EUR)
    put(BS, f"{c}22", f"={c}18+{c}19+{c}20+{c}21", "out", EUR)
BS.conditional_formatting.add("B23:F23", CellIsRule(operator="notEqual", formula=["0"], fill=FILL_FLAG))
widths(BS, {"A": 50, "B": 14, "C": 14, "D": 14, "E": 14, "F": 14})

# ---------------------------------------------------------------- Working Capital
WC = ws_["Working Capital"]
title(WC, "Capital circulante y desajuste temporal de caja",
      "Parte 1: saldos anuales. Parte 2: ciclo de caja mensual de un pedido (anticipo a Cuba → producción → transporte → frontera → ventas → cobro).")
put(WC, "A3", "=\"Escenario: \"&Assumptions!$D$4", "out")
statement(WC, 5, [
    ("Existencias", "inv_value", EUR, "calc"),
    ("IEE anticipado sobre existencias", "iee_prepaid", EUR, "calc"),
    ("Anticipos a fábrica", "prepay", EUR, "calc"),
    ("Clientes", "ar", EUR, "calc"),
    ("Hacienda deudora IVA", "vat_rec", EUR, "calc"),
    ("(-) Proveedores", "ap", EUR, "calc"),
    ("(-) Hacienda acreedora IVA", "vat_pay", EUR, "calc"),
    ("CAPITAL CIRCULANTE NETO", "nwc", EUR, "out"),
    ("Variación", "d_nwc", EUR, "calc"),
])
r0 = 18
section(WC, r0, "Parte 2 — Ciclo de caja de un pedido (meses desde el pedido)", 27)
put(WC, f"A{r0+1}", "Unidades del pedido")
put(WC, f"B{r0+1}", 3000, "input", NUM)
put(WC, f"A{r0+2}", "Meses de venta del lote tras llegada (incl. 1 mes de publicación de PVP en BOE)")
put(WC, f"B{r0+2}", 12, "input", NUM)
put(WC, f"A{r0+3}", "Incluir gastos fijos mensuales (1=sí)")
put(WC, f"B{r0+3}", 1, "input", "0")
put(WC, f"A{r0+4}", "Plazo pedido→llegada (meses, escenario)")
put(WC, f"B{r0+4}", f"={S['lead_time_sel']}", "calc", NUM)
put(WC, f"A{r0+5}", "Retraso de cobro (meses) = DSO/30")
put(WC, f"B{r0+5}", f"=ROUND({G['dso']}/30,0)", "calc", NUM)
U_, SP, FX_, LT, DL = (f"$B${r0+1}", f"$B${r0+2}", f"$B${r0+3}", f"$B${r0+4}", f"$B${r0+5}")
mr = r0 + 7
put(WC, f"A{mr}", "Mes")
for m in range(25):
    cc = WC.cell(row=mr, column=2 + m, value=m)
    cc.font, cc.fill, cc.alignment = F_H, FILL_H, CENTER
fob_pack = f"({SELREF('fob_eur',1)}+{SELREF('pack',1)}+{S['qc_sel']})"
flows = [
    ("Anticipo a fábrica", f"=IF({{c}}${mr}=0,-{G['advance']}*{U_}*{fob_pack},0)"),
    ("Saldo a fábrica antes de embarque", f"=IF({{c}}${mr}={LT}-1,-(1-{G['advance']})*{U_}*{fob_pack},0)"),
    ("Flete y seguro", f"=IF({{c}}${mr}={LT}-1,-{U_}*({SELREF('cif_u',1)}-{fob_pack}),0)"),
    ("Arancel + despacho en frontera", f"=IF({{c}}${mr}={LT},-({U_}*{SELREF('duty_u',1)}+{G['customs_fee']}+{U_}*({G['fiscal_marks']}+{S['torcedor_fee_sel']})),0)"),
    ("IEE (en frontera si se paga al importar; si no, al vender)", f"=IF({G['iee_at_import']}=1,IF({{c}}${mr}={LT},-{U_}*{SELREF('iee_u',1)},0),IF(AND({{c}}${mr}>{LT},{{c}}${mr}<={LT}+{SP}),-{U_}/{SP}*{SELREF('iee_u',1)},0))"),
    ("IVA de importación pagado", f"=IF({{c}}${mr}={LT},-{U_}*({SELREF('cif_u',1)}+{SELREF('duty_u',1)}+{G['iee_at_import']}*{SELREF('iee_u',1)})*{G['vat']},0)"),
    ("Devolución/compensación IVA importación (+3 meses)", f"=IF({{c}}${mr}={LT}+3,{U_}*({SELREF('cif_u',1)}+{SELREF('duty_u',1)}+{G['iee_at_import']}*{SELREF('iee_u',1)})*{G['vat']},0)"),
    ("Cobros del distribuidor (netos de IVA repercutido)", f"=IF(AND({{c}}${mr}>{LT}+1+{DL},{{c}}${mr}<={LT}+1+{DL}+{SP}-1),{U_}/{SP}/{SELREF('k',1)}*({SELREF('net_rev',2)}/MAX(1,{SELREF('sold',2)})),0)"),
    ("Gastos fijos mensuales", f"=IF({FX_}=1,-{SELREF('opex',1)}/12,0)"),
]
for n, (lab, f) in enumerate(flows):
    rr = mr + 1 + n
    WC.cell(row=rr, column=1, value=lab)
    for m in range(25):
        c = L(2 + m)
        put(WC, f"{c}{rr}", f.replace("{c}", c), "calc", NUM)
tot = mr + 1 + len(flows)
WC.cell(row=tot, column=1, value="Flujo neto del mes").font = F_OUT
WC.cell(row=tot + 1, column=1, value="Caja acumulada del pedido").font = F_OUT
for m in range(25):
    c = L(2 + m)
    put(WC, f"{c}{tot}", f"=SUM({c}{mr+1}:{c}{tot-1})", "out", NUM)
    put(WC, f"{c}{tot+1}", f"={c}{tot}" if m == 0 else f"={L(1+m)}{tot+1}+{c}{tot}", "out", NUM)
put(WC, f"A{tot+3}", "Necesidad máxima de caja del ciclo (−mínimo acumulado)")
put(WC, f"B{tot+3}", f"=-MIN(0,MIN(B{tot+1}:Z{tot+1}))", "out", EUR)
put(WC, f"A{tot+4}", "Mes en que se alcanza")
put(WC, f"B{tot+4}", f"=MATCH(MIN(B{tot+1}:Z{tot+1}),B{tot+1}:Z{tot+1},0)-1", "out", "0")
put(WC, f"A{tot+5}", "Meses entre el primer pago a Cuba y el primer cobro")
put(WC, f"B{tot+5}", f"={LT}+2+{DL}", "out", "0")
put(WC, f"A{tot+6}", "Nota: los cobros usan el ingreso neto unitario de 2029 del escenario; el ciclo ilustra el desajuste anticipo→cobro, no sustituye al modelo anual.", "note")
WC_PEAK = f"{q('Working Capital')}!$B${tot+3}"
widths(WC, {"A": 60})
for m in range(25):
    WC.column_dimensions[L(2 + m)].width = 10
# Gráfico caja acumulada
ch = LineChart()
ch.title = "Caja acumulada de un pedido (€)"
ch.height, ch.width = 7, 22
ch.add_data(Reference(WC, min_col=2, max_col=26, min_row=tot + 1), from_rows=True, titles_from_data=False)
ch.set_categories(Reference(WC, min_col=2, max_col=26, min_row=mr))
ch.legend = None
WC.add_chart(ch, f"A{tot+8}")

# ---------------------------------------------------------------- Break-even
BE = ws_["Break-even"]
title(BE, "Punto muerto (escenario seleccionado)", "Unidades/año necesarias para cubrir gastos fijos con la contribución unitaria del año.")
put(BE, "A3", "=\"Escenario: \"&Assumptions!$D$4", "out")
statement(BE, 5, [
    ("Contribución por unidad vendida (€)", "unit_contrib", EUR2, "calc"),
    ("Gastos fijos + despachos (€)", None, EUR, "calc"),
    ("Punto muerto EBITDA (unidades)", "be_ebitda", NUM, "out"),
    ("Punto muerto EBIT (unidades)", "be_ebit", NUM, "out"),
    ("Unidades vendidas", "sold", NUM, "calc"),
    ("Cobertura (ventas / punto muerto EBITDA)", None, PCT, "out"),
    ("Cajas/año para punto muerto EBITDA", None, NUM, "calc"),
    ("Puntos de venta necesarios para el punto muerto", None, NUM, "calc"),
])
for rr, lab in [(7, "Gastos fijos + despachos (€)"), (11, "Cobertura (ventas / punto muerto EBITDA)"),
                (12, "Cajas/año para punto muerto EBITDA"), (13, "Puntos de venta necesarios para el punto muerto")]:
    BE.cell(row=rr, column=1, value=lab).font = F_OUT if rr == 11 else F_CALC
    for cc in range(1, 7):
        BE.cell(row=rr, column=cc).fill = PatternFill(fill_type=None)
for i in range(NY):
    c = L(2 + i)
    put(BE, f"{c}7", f"={SELREF('opex',i)}+{G['shipments']}*{G['customs_fee']}", "calc", EUR)
    put(BE, f"{c}11", f"=IFERROR({c}10/{c}8,0)", "out", PCT)
    put(BE, f"{c}12", f"=IFERROR({c}8/{G['units_per_box']},\"n/a\")", "calc", NUM)
    put(BE, f"{c}13", f"=IFERROR({c}12/{G['boxes_per_pos']},\"n/a\")", "calc", NUM)
widths(BE, {"A": 56, "B": 14, "C": 14, "D": 14, "E": 14, "F": 14})

# ---------------------------------------------------------------- Investment Returns
IR = ws_["Investment Returns"]
title(IR, "Rentabilidad de la inversión (escenario seleccionado)",
      "FCF del proyecto sin financiación; valor terminal = EBITDA 2031 × múltiplo (si positivo). Descuento a fin de año.")
put(IR, "A3", "=\"Escenario: \"&Assumptions!$D$4", "out")
statement(IR, 5, [
    ("EBITDA", "ebitda", EUR, "calc"),
    ("(-) Impuestos", "tax", EUR, "calc"),
    ("(-) Variación circulante", "d_nwc", EUR, "calc"),
    ("(-) CAPEX", "capex", EUR, "calc"),
    ("FCF del proyecto", "fcf", EUR, "out"),
    ("Valor terminal", "tv", EUR, "calc"),
    ("FCF + valor terminal", "fcf_tv", EUR, "out"),
    ("Factor de descuento", "df", "0.0000", "calc"),
    ("Valor actual", "pv", EUR, "calc"),
    ("FCF acumulado", "cum_fcf", EUR, "calc"),
    ("NOPAT", "nopat", EUR, "calc"),
    ("Capital invertido", "invested", EUR, "calc"),
    ("ROIC", "roic", PCT, "out"),
])
put(IR, "A21", "VAN al coste de capital (€)")
put(IR, "B21", "=SUM(B14:F14)", "out", EUR)
put(IR, "A22", "TIR del proyecto")
put(IR, "B22", '=IFERROR(IRR(B12:F12),"no calculable")', "out", PCT)
put(IR, "A23", "TIR sin valor terminal")
put(IR, "B23", '=IFERROR(IRR(B10:F10),"no calculable")', "out", PCT)
put(IR, "A24", "Payback (año en que el FCF acumulado pasa a positivo)")
pbr = ROW["payback_flag"]
put(IR, "B24", f'=IFERROR(1/(1/MIN(INDEX(Engine!$C${pbr}:$Y${pbr},1,({SEL}-1)*6+1):INDEX(Engine!$C${pbr}:$Y${pbr},1,({SEL}-1)*6+5))),"no se recupera en 2027-2031")', "out", "0")
put(IR, "A25", "Inversión máxima acumulada (mínimo del FCF acumulado)")
put(IR, "B25", "=MIN(B15:F15)", "out", EUR)
put(IR, "A26", "Múltiplo de la inversión (VA flujos positivos / VA flujos negativos)")
put(IR, "B26", '=IFERROR(SUMIF(B14:F14,">0")/-SUMIF(B14:F14,"<0"),"n/a")', "out", X)
put(IR, "A27", "Valor atribuible a Arnaldo por su participación (equity × VAN positivo)")
put(IR, "B27", f"=MAX(0,B21)*{G['arnaldo_equity']}", "out", EUR)
put(IR, "C27", "=\"Royalties 2027-2031: \"&TEXT(" + "+".join(SELREF('royalty', i) for i in range(NY)) + ",\"#,##0\")&\" €\"", "note")
put(IR, "A29", "Advertencia: con 5 años de horizonte y valor terminal por múltiplo, la TIR es muy sensible al EBITDA de 2031. Ver 'Scenarios' y 'Sensitivity Analysis'.", "note")
widths(IR, {"A": 58, "B": 14, "C": 14, "D": 14, "E": 14, "F": 14})

# ---------------------------------------------------------------- Scenarios
SC = ws_["Scenarios"]
title(SC, "Comparativa de escenarios y adecuación del capital", "Todas las cifras leen directamente del motor (mismas fórmulas para los 4 escenarios).")
header(SC, 4, ["Métrica"] + SCEN)
metrics = [
    ("Unidades vendidas 2027–2031", lambda s: f"=SUM({ENGROW('sold',s)})", NUM),
    ("Unidades vendidas 2031", lambda s: f"={ENG('sold',s,4)}", NUM),
    ("Ingresos de la marca 2027–2031 (antes de IEE)", lambda s: f"=SUM({ENGROW('net_rev',s)})", EUR),
    ("Ingresos netos de impuestos 2031", lambda s: f"={ENG('net_rev_iee',s,4)}", EUR),
    ("EBITDA 2028", lambda s: f"={ENG('ebitda',s,1)}", EUR),
    ("EBITDA 2029", lambda s: f"={ENG('ebitda',s,2)}", EUR),
    ("EBITDA 2030", lambda s: f"={ENG('ebitda',s,3)}", EUR),
    ("EBITDA 2031", lambda s: f"={ENG('ebitda',s,4)}", EUR),
    ("EBITDA acumulado 2027–2031", lambda s: f"=SUM({ENGROW('ebitda',s)})", EUR),
    ("Resultado neto acumulado", lambda s: f"=SUM({ENGROW('net_income',s)})", EUR),
    ("FCF acumulado 2027–2031", lambda s: f"={ENG('cum_fcf',s,4)}", EUR),
    ("Necesidad máxima de financiación (capital + deuda)", lambda s: f"=MAX(0,{G['min_cash']}-MIN(0,MIN({ENGROW('cum_fcf',s)})))", EUR),
    ("Financiación ajena máxima con el capital de 'Assumptions'", lambda s: f"=MAX({ENGROW('debt',s)})", EUR),
    ("VAN (tasa de 'Assumptions')", lambda s: f"=SUM({ENGROW('pv',s)})", EUR),
    ("TIR del proyecto (con valor terminal)", lambda s: f"=IFERROR(IRR({ENGROW('fcf_tv',s)}),\"n/c\")", PCT),
    ("Payback (año)", lambda s: f"=IFERROR(1/(1/MIN({ENGROW('payback_flag',s)})),\"> 2031\")", "0"),
    ("Punto muerto EBITDA 2030 (unidades)", lambda s: f"={ENG('be_ebitda',s,3)}", NUM),
    ("Contribución unitaria 2030 (€/u)", lambda s: f"={ENG('unit_contrib',s,3)}", EUR2),
    ("ROIC 2031", lambda s: f"={ENG('roic',s,4)}", PCT),
    ("Control de balance (suma de descuadres)", lambda s: f"=SUMPRODUCT(ABS({ENGROW('check',s)}))", NUM2),
]
MROW = {}
for n, (lab, fn, fmt) in enumerate(metrics):
    rr = 5 + n
    SC.cell(row=rr, column=1, value=lab)
    MROW[lab] = rr
    for s in range(4):
        put(SC, f"{L(2+s)}{rr}", fn(s), "out" if n in (8, 11, 13, 14) else "calc", fmt)
need_row = MROW["Necesidad máxima de financiación (capital + deuda)"]
r = 5 + len(metrics) + 2
section(SC, r, "Adecuación del capital inicial (margen de seguridad sobre la necesidad máxima)", 6)
put(SC, f"A{r+1}", "Colchón de seguridad exigido sobre la necesidad máxima")
put(SC, f"B{r+1}", 0.25, "input", PCT)
cush = f"$B${r+1}"
header(SC, r + 2, ["Capital disponible"] + SCEN)
cap_levels = [15000, 30000, 50000, 75000, 100000]
for n, cap in enumerate(cap_levels):
    rr = r + 3 + n
    put(SC, f"A{rr}", cap, "input", EUR)
    for s in range(4):
        c = L(2 + s)
        put(SC, f"{c}{rr}", f'=IF($A{rr}>={c}${need_row}*(1+{cush}),"Viable",IF($A{rr}>={c}${need_row},"Ajustado","Insuficiente"))', "calc")
cap_end = r + 3 + len(cap_levels)
for s in range(4):
    c = L(2 + s)
    SC.conditional_formatting.add(f"{c}{r+3}:{c}{cap_end-1}", CellIsRule(operator="equal", formula=['"Insuficiente"'], fill=FILL_FLAG))
    SC.conditional_formatting.add(f"{c}{r+3}:{c}{cap_end-1}", CellIsRule(operator="equal", formula=['"Viable"'], fill=FILL_OUT))
r = cap_end + 1
section(SC, r, "Tamaño del lote inicial (escenario seleccionado, valores 2028)", 8)
header(SC, r + 1, ["Unidades del lote", "Pago a fábrica (€)", "Coste en destino total (€)", "IEE si se paga al importar (€)", "Caja inmovilizada (€)", "Meses de venta a demanda 2029", "Riesgo de envejecimiento"])
for n, lot in enumerate([1000, 3000, 5000, 10000, 20000]):
    rr = r + 2 + n
    put(SC, f"A{rr}", lot, "input", NUM)
    put(SC, f"B{rr}", f"=A{rr}*({SELREF('fob_eur',1)}+{SELREF('pack',1)}+{S['qc_sel']})", "calc", EUR)
    put(SC, f"C{rr}", f"=A{rr}*{SELREF('landed_u',1)}+{G['customs_fee']}", "calc", EUR)
    put(SC, f"D{rr}", f"=A{rr}*{SELREF('iee_u',1)}*{G['iee_at_import']}", "calc", EUR)
    put(SC, f"E{rr}", f"=C{rr}+D{rr}", "out", EUR)
    put(SC, f"F{rr}", f"=IFERROR(A{rr}/({SELREF('demand',2)}/12),\"n/a\")", "calc", "0.0")
    put(SC, f"G{rr}", f'=IF(F{rr}="n/a","n/a",IF(F{rr}>24,"Alto",IF(F{rr}>12,"Medio","Bajo")))', "calc")
widths(SC, {"A": 56, "B": 16, "C": 18, "D": 18, "E": 16, "F": 16, "G": 16})
# Gráfico EBITDA por escenario
ch = BarChart()
ch.type = "col"
ch.title = "EBITDA por escenario 2028–2031 (€)"
ch.height, ch.width = 8, 20
data = Reference(SC, min_col=2, max_col=5, min_row=MROW["EBITDA 2028"], max_row=MROW["EBITDA 2031"])
ch.add_data(data, titles_from_data=False)
from openpyxl.chart.series import SeriesLabel
for s, ser in enumerate(ch.series):
    ser.tx = SeriesLabel(v=SCEN[s])
ch.set_categories(Reference(SC, min_col=1, min_row=MROW["EBITDA 2028"], max_row=MROW["EBITDA 2031"]))
SC.add_chart(ch, "I4")

# ---------------------------------------------------------------- Sensitivity
SA = ws_["Sensitivity Analysis"]
title(SA, "Análisis de sensibilidad — EBITDA 2030 (año estabilizado, escenario seleccionado)",
      "Forma cerrada equivalente al motor para 2030: EBITDA = U × contribución(PVP, FOB, FX, tarifa, merma) − gastos fijos − despachos.")
put(SA, "A4", "Base: unidades 2030")
put(SA, "B4", "=" + SELREF("sold", 3), "calc", NUM)
put(SA, "A5", "Base: PVP")
put(SA, "B5", "=" + SELREF("pvp", 3), "calc", EUR2)
put(SA, "A6", "Base: FOB USD/u")
put(SA, "B6", "=" + S["fob_usd_sel"], "calc", NUM2)
put(SA, "A7", "Base: tipo de cambio USD/€")
put(SA, "B7", "=" + G["fx"], "calc", NUM2)
put(SA, "A8", "Base: tarifa distribuidor")
put(SA, "B8", "=" + S["dist_fee_sel"], "calc", PCT)
put(SA, "A9", "Base: merma")
put(SA, "B9", "=" + S["merma_sel"], "calc", PCT)
put(SA, "A10", "Coste no-FOB en destino por unidad (excl. FOB y su arancel/seguro)")
put(SA, "B10", f"={SELREF('landed_u',3)}-{SELREF('fob_eur',3)}*(1+{G['insurance']})*(1+{G['duty']})", "calc", EUR2)
put(SA, "A11", "Gastos fijos 2030 + despachos (motor)")
put(SA, "B11", f"={SELREF('opex',3)}+{SELREF('customs_fees',3)}", "calc", EUR)
put(SA, "A12", "EBITDA 2030 (motor)")
put(SA, "B12", "=" + SELREF("ebitda", 3), "out", EUR)


def ebitda_expr(U="$B$4", P="$B$5", FOB="$B$6", FX="$B$7", D="$B$8", M="$B$9"):
    k = f"(1/(1-{M})+{S['muestras_sel']})"
    landed = f"({FOB}/{FX}*(1+{G['insurance']})*(1+{G['duty']})+$B$10)"
    iee = f"MAX({G['iee_rate']}*{P},{G['iee_min']}/1000)"
    contrib = f"(({P}/(1+{G['vat']})-{P}*({G['retail_margin']}+{D}))*(1-{G['royalty']})-{k}*({iee}+{landed}))"
    return f"{U}*{contrib}-$B$11"


put(SA, "A13", "EBITDA 2030 (forma cerrada) — control")
put(SA, "B13", "=" + ebitda_expr(), "out", EUR)
put(SA, "C13", "=ROUND(B13-B12,0)", "calc", NUM)
put(SA, "D13", "← diferencia (debe ser 0)", "note")
r = 15
section(SA, r, "Tornado: impacto en EBITDA 2030 de variaciones individuales", 6)
header(SA, r + 1, ["Variable", "Caso desfavorable", "Caso favorable", "EBITDA desfavorable", "EBITDA favorable", "Rango"])
tor = [("Volumen (unidades)", -0.2, 0.2, "U"), ("PVP", -0.1, 0.1, "P"), ("Coste FOB", 0.2, -0.2, "FOB"),
       ("Tipo de cambio USD/€", -0.1, 0.1, "FX"), ("Tarifa distribuidor (pp)", 0.02, -0.02, "D"), ("Merma (pp)", 0.03, -0.02, "M")]
for n, (lab, lo, hi, var) in enumerate(tor):
    rr = r + 2 + n
    put(SA, f"A{rr}", lab)
    put(SA, f"B{rr}", lo, "input", PCT)
    put(SA, f"C{rr}", hi, "input", PCT)
    for colx, cell in (("D", f"$B{rr}"), ("E", f"$C{rr}")):
        args = {}
        if var in ("U", "P", "FOB", "FX"):
            base = {"U": "$B$4", "P": "$B$5", "FOB": "$B$6", "FX": "$B$7"}[var]
            args[var] = f"({base}*(1+{cell}))"
        else:
            base = {"D": "$B$8", "M": "$B$9"}[var]
            args[var] = f"({base}+{cell})"
        put(SA, f"{colx}{rr}", "=" + ebitda_expr(**args), "calc", EUR)
    put(SA, f"F{rr}", f"=ABS(E{rr}-D{rr})", "out", EUR)
r = r + 2 + len(tor) + 1
section(SA, r, "Tabla bidimensional: EBITDA 2030 según PVP (filas) y FOB USD/u (columnas)", 9)
fobs = [4, 5, 6, 7, 8, 9, 10]
pvps = [12, 15, 18, 20, 25, 30, 35]
header(SA, r + 1, ["PVP \\ FOB"] + [f"{f} USD" for f in fobs])
for j, f in enumerate(fobs):
    put(SA, f"{L(2+j)}{r+2}", f, "input", NUM2)
SA.cell(row=r + 2, column=1, value="FOB (input)")
for n, p in enumerate(pvps):
    rr = r + 3 + n
    put(SA, f"A{rr}", p, "input", EUR2)
    for j in range(len(fobs)):
        c = L(2 + j)
        put(SA, f"{c}{rr}", "=" + ebitda_expr(P=f"$A{rr}", FOB=f"{c}${r+2}"), "calc", EUR)
    SA.conditional_formatting.add(f"B{rr}:H{rr}", CellIsRule(operator="lessThan", formula=["0"], fill=FILL_FLAG))
r = r + 3 + len(pvps) + 1
section(SA, r, "Tabla bidimensional: EBITDA 2030 según volumen (filas) y tarifa del distribuidor (columnas)", 9)
dfs = [0.03, 0.04, 0.05, 0.06, 0.08, 0.10]
vols = [2000, 4000, 6000, 8000, 10000, 15000, 20000]
header(SA, r + 1, ["Unidades \\ tarifa"] + [f"{int(d*100)}%" for d in dfs])
for j, d in enumerate(dfs):
    put(SA, f"{L(2+j)}{r+2}", d, "input", PCT)
for n, v in enumerate(vols):
    rr = r + 3 + n
    put(SA, f"A{rr}", v, "input", NUM)
    for j in range(len(dfs)):
        c = L(2 + j)
        put(SA, f"{c}{rr}", "=" + ebitda_expr(U=f"$A{rr}", D=f"{c}${r+2}"), "calc", EUR)
    SA.conditional_formatting.add(f"B{rr}:G{rr}", CellIsRule(operator="lessThan", formula=["0"], fill=FILL_FLAG))
r = r + 3 + len(vols) + 1
section(SA, r, "Retrasos de suministro y coste de financiación (aproximación)", 6)
put(SA, f"A{r+1}", "Gastos fijos mensuales 2028")
put(SA, f"B{r+1}", f"={SELREF('opex',1)}/12", "calc", EUR)
header(SA, r + 2, ["Retraso adicional (meses)", "Caja adicional necesaria (€)", "Coste financiero anual al 5%", "al 7%", "al 10%", "al 12%"])
for n, dm in enumerate([0, 3, 6, 9, 12]):
    rr = r + 3 + n
    put(SA, f"A{rr}", dm, "input", NUM)
    put(SA, f"B{rr}", f"=A{rr}*$B${r+1}", "calc", EUR)
    for j, rate in enumerate([0.05, 0.07, 0.10, 0.12]):
        put(SA, f"{L(3+j)}{rr}", f"=({SELREF('debt',1)}+B{rr})*{rate}", "calc", EUR)
put(SA, f"A{r+9}", "Lectura: cada mes de retraso consume gastos fijos sin ingresos; el anticipo pagado a Cuba ya está inmovilizado. La columna de coste financiero incluye la deuda del escenario en 2028.", "note")
widths(SA, {"A": 56, "B": 16, "C": 16, "D": 16, "E": 16, "F": 16, "G": 14, "H": 14})

# ---------------------------------------------------------------- Dashboard
D = ws_["Dashboard"]
title(D, "ARNALDO VALLÉS — Modelo financiero 2027–2031 · Dashboard ejecutivo",
      "Proyecto A (marca propia fabricada en Cuba). Condicionado a la confirmación jurídica de la Fase 0. Fecha de corte 10/10/2026.")
put(D, "A4", "Escenario activo (cambiar en Assumptions!C4)")
put(D, "B4", "=Assumptions!$D$4", "out")
kpis = [
    ("Unidades vendidas 2031", "=" + SELREF("sold", 4), NUM),
    ("Ingresos de la marca 2031 (antes de IEE)", "=" + SELREF("net_rev", 4), EUR),
    ("EBITDA 2031", "=" + SELREF("ebitda", 4), EUR),
    ("EBITDA acumulado 2027–2031", f"=SUM({''.join([])}{'+'.join(SELREF('ebitda', i) for i in range(NY))})", EUR),
    ("Resultado neto acumulado", "=" + "+".join(SELREF("net_income", i) for i in range(NY)), EUR),
    ("Necesidad máxima de financiación (capital + deuda)", f"=INDEX(Scenarios!$B${need_row}:$E${need_row},1,{SEL})", EUR),
    ("Capital aportado (Assumptions)", "=" + G["capital"], EUR),
    ("Financiación ajena máxima requerida", "=" + "MAX(" + ",".join(SELREF("debt", i) for i in range(NY)) + ")", EUR),
    ("VAN", "='Investment Returns'!B21", EUR),
    ("TIR del proyecto (con valor terminal)", "='Investment Returns'!B22", PCT),
    ("Payback", "='Investment Returns'!B24", "0"),
    ("Punto muerto EBITDA 2030 (unidades/año)", "=" + SELREF("be_ebitda", 3), NUM),
    ("Contribución unitaria 2030 (€/u)", "=" + SELREF("unit_contrib", 3), EUR2),
    ("Necesidad de caja de un pedido piloto (Working Capital)", "=" + WC_PEAK, EUR),
    ("Control de balance (0 = cuadra)", "=SUMPRODUCT(ABS(Engine!$C$%d:$Y$%d))" % (ROW["check"], ROW["check"]), NUM2),
]
header(D, 6, ["Indicador", "Valor"])
for n, (lab, f, fmt) in enumerate(kpis):
    rr = 7 + n
    put(D, f"A{rr}", lab)
    put(D, f"B{rr}", f, "out", fmt)
rr = 7 + len(kpis) + 1
put(D, f"A{rr}", "Veredicto financiero automático")
put(D, f"B{rr}", f'=IF(B{7+7}>0,"El capital no cubre la necesidad: requiere financiación adicional",IF(B{7+2}<=0,"EBITDA 2031 negativo: modelo no sostenible","Autofinanciable con el capital indicado"))', "flag")
put(D, f"A{rr+1}", "Veredicto jurídico (no calculable)")
put(D, f"B{rr+1}", "BLOQUEADO hasta confirmar por escrito la vía legal de fabricación y exportación en Cuba (Fase 0)", "flag")
rr += 3
header(D, rr, ["Comparativa", *SCEN])
comp = ["Unidades vendidas 2031", "EBITDA 2031", "EBITDA acumulado 2027–2031", "Necesidad máxima de financiación (capital + deuda)",
        "VAN (tasa de 'Assumptions')", "TIR del proyecto (con valor terminal)", "Payback (año)", "Punto muerto EBITDA 2030 (unidades)"]
for n, lab in enumerate(comp):
    r2 = rr + 1 + n
    put(D, f"A{r2}", lab)
    src = MROW[lab]
    fmt = PCT if "TIR" in lab else ("0" if "Payback" in lab else (NUM if "Unidades" in lab or "unidades" in lab else EUR))
    for s in range(4):
        put(D, f"{L(2+s)}{r2}", f"=Scenarios!{L(2+s)}{src}", "calc", fmt)
widths(D, {"A": 58, "B": 22, "C": 16, "D": 16, "E": 16})
ch = BarChart()
ch.type = "col"
ch.title = "EBITDA del escenario activo (€)"
ch.height, ch.width = 7, 16
# serie auxiliar
aux = rr + 11
put(D, f"A{aux}", "Serie gráfica: EBITDA escenario activo", "note")
for i in range(NY):
    put(D, f"{L(2+i)}{aux-1}", YEARS[i], "calc", "0")
    put(D, f"{L(2+i)}{aux}", "=" + SELREF("ebitda", i), "calc", EUR)
ch.add_data(Reference(D, min_col=2, max_col=6, min_row=aux), from_rows=True, titles_from_data=False)
ch.set_categories(Reference(D, min_col=2, max_col=6, min_row=aux - 1))
ch.legend = None
D.add_chart(ch, "D4")

# ---------------------------------------------------------------- Risk Register (resumen)
RR = ws_["Risk Register"]
title(RR, "Riesgos con impacto en el modelo (resumen; registro completo en 07_OPERATIONS/Risk Register.xlsx)", "")
risks = [
    ("R01", "No existe vía legal para fabricar/exportar una marca de titularidad extranjera en Cuba", "Crítico", "Todo el proyecto A", "Fase 0: confirmación escrita del exportador y del organismo competente antes de invertir"),
    ("R02", "Habanos S.A./Cubatabaco exige titularidad o control de la marca", "Crítico", "Valor de la marca", "No ceder la marca; si es condición, NO-GO o licencia con opción de recompra"),
    ("R03", "Coste FOB superior al supuesto (sin cotización)", "Alto", "Contribución unitaria", "Sensibilidad PVP×FOB; FOB máximo en 'Unit Economics'"),
    ("R04", "Retrasos de producción/transporte (anticipo inmovilizado)", "Alto", "Caja", "Lotes pequeños, hitos de pago, seguro"),
    ("R05", "Imposibilidad de pagar a Cuba por canal bancario", "Alto", "Operativa", "Validar banco antes de firmar; carta de crédito"),
    ("R06", "Distribuidor no acepta la marca o exige tarifa mayor", "Alto", "Volumen/margen", "Cartas de interés previas; tabla volumen×tarifa"),
    ("R07", "Cambio normativo UE (fiscalidad, TPD) 2028+", "Medio", "Precio/costes", "Seguimiento; el mínimo propuesto afecta poco a PVP ≥12 €"),
    ("R08", "Rotación lenta en expendedurías (envejecimiento, merma)", "Medio", "Circulante", "Cobertura 6 meses; humidificación en 3PL"),
    ("R09", "Calidad inconsistente entre lotes", "Medio", "Marca", "Ficha técnica, catas ciegas, rechazo contractual"),
    ("R10", "Bloqueo de suministro (escenario Estrés)", "Alto", "Continuidad", "Plan de contingencia; marca transferible a otro origen (con cambio de producto)"),
]
table(RR, 4, ["ID", "Riesgo", "Severidad", "Variable afectada", "Mitigación"], risks, [6, 70, 10, 22, 70])

# ---------------------------------------------------------------- Sources & Validation
SV = ws_["Sources & Validation"]
title(SV, "Fuentes y validación de supuestos", "Fecha de consulta: 10/10/2026. Acceso mediante motor de búsqueda (descarga directa bloqueada en el entorno).")
rows_src = [(s[0], s[1], s[3], s[4], s[5], s[2]) for s in SOURCES]
table(SV, 4, ["ID", "Fuente", "Tipo", "Dato utilizado", "Estado", "URL"], rows_src, [6, 50, 18, 70, 22, 60])
r = 6 + len(rows_src)
section(SV, r, "Controles automáticos del modelo", 6)
checks = [
    ("Balance cuadra en los 4 escenarios", "=IF(SUMPRODUCT(ABS(Engine!$C$%d:$Y$%d))<0.5,\"OK\",\"ERROR\")" % (ROW["check"], ROW["check"])),
    ("Reparto del PVP suma 100% (Import & Taxes)", f"=IF(ABS({q('Import & Taxes')}!D{p0+len(split)}-1)<0.0001,\"OK\",\"ERROR\")"),
    ("Caja final conciliada con variación (Cash Flow)", "=IF(SUMPRODUCT(ABS('Cash Flow'!B19:F19))<0.5,\"OK\",\"ERROR\")"),
    ("Sensibilidad: forma cerrada = motor (EBITDA 2030)", "=IF(ABS('Sensitivity Analysis'!C13)<1,\"OK\",\"ERROR\")"),
    ("Stock final nunca negativo", "=IF(MIN(Engine!$C$%d:$Y$%d)>=-0.01,\"OK\",\"ERROR\")" % (ROW["inv_end"], ROW["inv_end"])),
]
for n, (lab, f) in enumerate(checks):
    put(SV, f"A{r+1+n}", lab)
    SV.merge_cells(f"A{r+1+n}:C{r+1+n}")
    put(SV, f"D{r+1+n}", f, "out")

# Orden de pestañas: Dashboard primero
wb.active = 0
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT, "engine rows", ROW["check"], "need_row", need_row)
