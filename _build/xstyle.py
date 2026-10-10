"""Estilos compartidos para todos los XLSX del proyecto.

Convención visual (igual en todos los libros):
  - INPUT  : fuente azul sobre fondo crema  -> supuesto editable
  - CALC   : fuente negra                   -> fórmula
  - OUTPUT : negrita sobre fondo verde suave -> resultado clave
  - FLAG   : fondo rojo suave               -> alerta / bloqueo
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

INK = "1B1A17"
TOBACCO = "5A3E2B"
VERMILION = "A23B2A"
GOLD = "B08D57"
IVORY = "F2EDE3"

F_TITLE = Font(name="Calibri", size=16, bold=True, color=INK)
F_SUB = Font(name="Calibri", size=10, italic=True, color="666666")
F_H = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
F_SEC = Font(name="Calibri", size=11, bold=True, color=TOBACCO)
F_INPUT = Font(name="Calibri", size=10, color="1F4E9E")
F_CALC = Font(name="Calibri", size=10, color="000000")
F_OUT = Font(name="Calibri", size=10, bold=True, color="000000")
F_NOTE = Font(name="Calibri", size=9, italic=True, color="666666")
F_FLAG = Font(name="Calibri", size=10, bold=True, color="9C0006")

FILL_H = PatternFill("solid", fgColor=TOBACCO)
FILL_SEC = PatternFill("solid", fgColor="EFE7DA")
FILL_INPUT = PatternFill("solid", fgColor="FFF8E1")
FILL_OUT = PatternFill("solid", fgColor="E2F0D9")
FILL_FLAG = PatternFill("solid", fgColor="FADBD8")
FILL_GREY = PatternFill("solid", fgColor="F2F2F2")

THIN = Side(style="thin", color="BFBFBF")
B_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
B_BOT = Border(bottom=Side(style="thin", color=TOBACCO))

EUR = '#,##0" €";[Red]-#,##0" €"'
EUR2 = '#,##0.00" €";[Red]-#,##0.00" €"'
NUM = '#,##0;[Red]-#,##0'
NUM2 = '#,##0.00;[Red]-#,##0.00'
PCT = '0.0%'
X = '0.00"x"'

WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def title(ws, text, sub=None, width_a=46):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F_SUB
    ws.column_dimensions["A"].width = width_a
    ws.sheet_view.showGridLines = False


def header(ws, row, values, col=1, fill=FILL_H, font=F_H):
    for i, v in enumerate(values):
        c = ws.cell(row=row, column=col + i, value=v)
        c.fill = fill
        c.font = font
        c.alignment = CENTER
        c.border = B_ALL


def section(ws, row, text, ncols=8):
    for i in range(1, ncols + 1):
        ws.cell(row=row, column=i).fill = FILL_SEC
    c = ws.cell(row=row, column=1, value=text)
    c.font = F_SEC


def put(ws, ref, value, kind="calc", fmt=None):
    c = ws[ref]
    c.value = value
    if kind == "input":
        c.font, c.fill = F_INPUT, FILL_INPUT
    elif kind == "out":
        c.font, c.fill = F_OUT, FILL_OUT
    elif kind == "flag":
        c.font, c.fill = F_FLAG, FILL_FLAG
    elif kind == "note":
        c.font = F_NOTE
    else:
        c.font = F_CALC
    if fmt:
        c.number_format = fmt
    if kind in ("input", "out", "calc", "flag"):
        c.border = B_ALL
    return c


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def legend(ws, row, col=1):
    items = [("Supuesto editable (input)", F_INPUT, FILL_INPUT),
             ("Fórmula", F_CALC, None),
             ("Resultado clave", F_OUT, FILL_OUT),
             ("Alerta / bloqueo", F_FLAG, FILL_FLAG)]
    for i, (t, f, fill) in enumerate(items):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font = f
        if fill:
            c.fill = fill
        c.border = B_ALL
        c.alignment = CENTER


def table(ws, start_row, headers, rows, col_widths=None, wrap=True, input_cols=()):
    """Escribe una tabla de texto/valores. input_cols: índices (0-based) que se marcan como input."""
    header(ws, start_row, headers)
    for r_i, r in enumerate(rows):
        for c_i, v in enumerate(r):
            c = ws.cell(row=start_row + 1 + r_i, column=1 + c_i, value=v)
            c.border = B_ALL
            c.alignment = WRAP if wrap else Alignment(vertical="top")
            if c_i in input_cols:
                c.font, c.fill = F_INPUT, FILL_INPUT
            else:
                c.font = F_CALC
    if col_widths:
        for i, w in enumerate(col_widths):
            ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    return start_row + 1 + len(rows)
